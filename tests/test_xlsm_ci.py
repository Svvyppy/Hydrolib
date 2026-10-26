"""Exercise CI with a local bare Git remote; never contacts GitHub."""
import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

TOOLS = Path(__file__).resolve().parents[1] / 'tools/xlsm_merge'
sys.path.insert(0, str(TOOLS))


class PipelineTests(unittest.TestCase):
    def git(self, *args, cwd=None):
        return subprocess.run(['git', *args], cwd=cwd or self.seed, check=True,
                              capture_output=True).stdout.decode().strip()

    def setUp(self):
        # Hard isolation: even a mistaken remote URL cannot use HTTPS/SSH.
        isolation = patch.dict(os.environ, {'GIT_ALLOW_PROTOCOL': 'file',
            'GIT_CONFIG_GLOBAL': os.devnull, 'GIT_CONFIG_SYSTEM': os.devnull,
            'GIT_TERMINAL_PROMPT': '0'})
        isolation.start()
        self.addCleanup(isolation.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.remote = self.root / 'remote.git'
        self.seed = self.root / 'seed'
        self.seed.mkdir()
        self.git('init', '--bare', str(self.remote))
        self.git('init', '-b', 'master')
        self.git('config', 'user.name', 'Test')
        self.git('config', 'user.email', 'test@example.com')
        self.book = self.seed / 'Database/Hydrolib_bd.xlsm'
        self.book.parent.mkdir()
        shutil.copytree(TOOLS, self.seed / 'tools/xlsm_merge', ignore=shutil.ignore_patterns('__pycache__'))
        (self.seed / '.gitattributes').write_text('Database/Hydrolib_bd.xlsm merge=hydrolib-xlsm\n')
        w = openpyxl.Workbook()
        s = w.active
        s.append(['Part Number', 'Value', 'Left', 'Right', 'Library Path', 'Footprint Path 1'])
        s.append(['A', 10, 'base', 'base'])
        w.save(self.book)
        self.ancestor = self.book.read_bytes()
        self.git('add', '.')
        self.git('commit', '-m', 'ancestor')
        self.git('switch', '-c', 'feature')
        s['B2'] = 20
        s['C2'] = 'ours independent'
        w.save(self.book)
        self.ours = self.book.read_bytes()
        self.git('add', '.')
        self.git('commit', '-m', 'ours')
        self.head_sha = self.git('rev-parse', 'HEAD')
        self.git('switch', 'master')
        s['B2'] = 30
        s['C2'] = 'base'
        s['D2'] = 'theirs independent'
        w.save(self.book)
        self.theirs = self.book.read_bytes()
        self.git('add', '.')
        self.git('commit', '-m', 'theirs')
        self.base_sha = self.git('rev-parse', 'HEAD')
        self.git('remote', 'add', 'origin', str(self.remote))
        self.git('push', 'origin', 'master', 'feature')
        self.checkout = self.root / 'checkout'
        self.git('clone', '-b', 'master', str(self.remote), str(self.checkout))
        spec = importlib.util.spec_from_file_location('pipeline_under_test', TOOLS / 'ci.py')
        self.ci = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.ci)
        self.artifacts = self.root / 'artifacts'
        self.artifacts.mkdir()
        self.ci.ROOT = self.checkout
        self.ci.ARTIFACTS = self.artifacts
        self.pr = {'state': 'open', 'base': {'sha': self.base_sha, 'ref': 'master'}, 'head': {
            'sha': self.head_sha, 'ref': 'feature', 'repo': {'full_name': 'test/library'}}}

    def invoke(self, document=None, apply=False, side='manual'):
        env = {'PR_NUMBER': '1', 'GITHUB_REPOSITORY': 'test/library', 'GH_TOKEN': 'test-token',
               'RESOLUTIONS_JSON': json.dumps(document) if document else '', 'CONFLICT_SIDE': side}
        with patch.dict(os.environ, env), patch.object(sys, 'argv', ['ci.py'] + (['--apply'] if apply else [])), \
                patch.object(self.ci.urllib.request, 'urlopen', side_effect=lambda *a, **k: BytesIO(json.dumps(self.pr).encode())):
            self.ci.main()

    def test_report_then_apply_selected_side_and_keep_independent_changes(self):
        with self.assertRaisesRegex(ValueError, 'unresolved'):
            self.invoke()
        report = json.loads((self.artifacts / 'report.json').read_text())
        self.assertEqual(len(report['conflicts']), 1)
        self.assertEqual(report['conflicts'][0]['ours'], 20)
        proposed = json.loads((self.artifacts / 'pr-changes.json').read_text())
        fields = {c['field']: c for c in proposed['changes']['modified']}
        self.assertEqual(fields['Value']['before'], 10)
        self.assertEqual(fields['Value']['after'], 20)
        self.assertEqual(fields['Left']['after'], 'ours independent')
        self.assertNotIn('Right', fields)

        for c in report['conflicts']:
            c['choice'] = 'theirs'
        self.invoke(report, apply=True)
        data = subprocess.run(['git', '--git-dir', str(self.remote), 'show', 'feature:Database/Hydrolib_bd.xlsm'],
                              check=True, capture_output=True).stdout
        s = openpyxl.load_workbook(BytesIO(data)).active
        self.assertEqual(s['B2'].value, 30)
        self.assertEqual(s['C2'].value, 'ours independent')
        self.assertEqual(s['D2'].value, 'theirs independent')
        parents = self.git('--git-dir', str(self.remote), 'log', '-1', '--format=%P', 'feature').split()
        self.assertEqual(parents, [self.head_sha, self.base_sha])

    def test_dropdown_choice_applies_without_json(self):
        self.invoke(apply=True, side='ours')
        data = subprocess.run(['git', '--git-dir', str(self.remote), 'show', 'feature:Database/Hydrolib_bd.xlsm'],
                              check=True, capture_output=True).stdout
        s = openpyxl.load_workbook(BytesIO(data)).active
        self.assertEqual(s['B2'].value, 20)
        self.assertEqual(s['C2'].value, 'ours independent')
        self.assertEqual(s['D2'].value, 'theirs independent')
        report = json.loads((self.artifacts / 'report.json').read_text())
        self.assertEqual(report['conflicts'][0]['choice'], 'ours')

    def fork_api(self, path, method='GET', payload=None):
        if path == 'pulls/1':
            return self.pr
        if method == 'GET':
            return self.created
        self.payload = payload
        self.created.append({'html_url': 'https://github.com/test/library/pull/2', 'number': 2})
        return self.created[0]

    def prepare_fork(self):
        self.pr['head']['repo']['full_name'] = 'contributor/library'
        self.git('--git-dir', str(self.remote), 'update-ref', 'refs/pull/1/head', self.head_sha)
        self.created = []

    def test_fork_creates_result_in_target_and_retry_reuses_pr(self):
        self.prepare_fork()
        with patch.object(self.ci, 'api', side_effect=self.fork_api):
            self.invoke(apply=True, side='ours')
            result = json.loads((self.artifacts / 'result.json').read_text())
            self.assertEqual(result['mode'], 'fork_result')
            self.assertEqual(self.payload['base'], 'master')
            self.assertEqual(self.payload['head'], result['branch'])
            self.assertIn('PR #1', self.payload['body'])
            self.assertTrue(result['branch'].startswith('xlsm-resolved/pr-1-'))
            self.assertEqual(self.git('--git-dir', str(self.remote), 'rev-parse', 'feature'), self.head_sha)
            self.assertEqual(self.git('--git-dir', str(self.remote), 'rev-parse', 'master'), self.base_sha)
            self.invoke(apply=True, side='ours')
            self.assertEqual(len(self.created), 1)
            self.assertEqual(json.loads((self.artifacts / 'result.json').read_text()), result)

    def test_fork_snapshot_change_prevents_result_push(self):
        self.prepare_fork()
        calls = []
        def api(path, **kwargs):
            calls.append(path)
            if len(calls) > 1:
                return {**self.pr, 'state': 'closed'}
            return self.pr
        with patch.object(self.ci, 'api', side_effect=api):
            with self.assertRaisesRegex(ValueError, 'изменился'):
                self.invoke(apply=True, side='ours')
        self.assertFalse(self.git('--git-dir', str(self.remote), 'for-each-ref', '--format=%(refname)', 'refs/heads/xlsm-resolved/'))

    def prepare_library_refs(self, add_files):
        self.git('switch', 'feature')
        w = openpyxl.load_workbook(self.book)
        w.active['E2'] = '=CONCATENATE("Schematic/",A2,".SchLib")'
        w.active['F2'] = 'Footprints/A.PcbLib'
        w.save(self.book)
        w.close()
        if add_files:
            for path in ('Schematic/A.SchLib', 'Footprints/A.PcbLib'):
                p = self.seed / path
                p.parent.mkdir(exist_ok=True)
                p.write_text('test library bytes')
        self.git('add', '.')
        self.git('commit', '-m', 'component library refs')
        self.head_sha = self.git('rev-parse', 'HEAD')
        self.pr['head']['sha'] = self.head_sha
        self.git('push', 'origin', 'feature')

    def test_new_library_files_in_same_pr_are_found(self):
        self.prepare_library_refs(add_files=True)
        self.invoke(apply=True, side='ours')
        report = json.loads((self.artifacts / 'validation.json').read_text())
        self.assertEqual(report['errors'], 0)
        self.assertEqual(report['phase'], 'published result')
        self.assertIn('Schematic/A.SchLib', self.ci.tracked_files('HEAD', cwd=self.seed))

    def test_missing_library_files_block_push(self):
        self.prepare_library_refs(add_files=False)
        with self.assertRaisesRegex(ValueError, 'Проверка компонентов'):
            self.invoke(apply=True, side='ours')
        self.assertEqual(self.git('--git-dir', str(self.remote), 'rev-parse', 'feature'), self.head_sha)
        report = json.loads((self.artifacts / 'validation.json').read_text())
        self.assertEqual(report['errors'], 2)
        self.assertEqual({i['code'] for i in report['issues']}, {'missing_file'})

    def test_invalid_component_blocks_apply_before_push(self):
        self.git('switch', 'feature')
        w = openpyxl.load_workbook(self.book)
        w.active.append([None, 'new invalid component'])
        w.save(self.book)
        w.close()
        self.git('add', '.')
        self.git('commit', '-m', 'invalid component')
        self.head_sha = self.git('rev-parse', 'HEAD')
        self.pr['head']['sha'] = self.head_sha
        self.git('push', 'origin', 'feature')
        with self.assertRaisesRegex(ValueError, 'Проверка компонентов'):
            self.invoke(apply=True, side='ours')
        self.assertEqual(self.git('--git-dir', str(self.remote), 'rev-parse', 'feature'), self.head_sha)
        report = json.loads((self.artifacts / 'validation.json').read_text())
        self.assertEqual(report['errors'], 1)
        self.assertEqual(report['issues'][0]['cell'], 'A3')
        self.assertTrue((self.artifacts / 'pr-changes.json').exists())
        self.assertFalse((self.artifacts / 'result.json').exists())

    def test_unresolved_apply_does_not_push(self):
        with self.assertRaisesRegex(ValueError, 'unresolved'):
            self.invoke(apply=True)
        actual = self.git('--git-dir', str(self.remote), 'rev-parse', 'feature')
        self.assertEqual(actual, self.head_sha)
        self.assertTrue((self.artifacts / 'report.json').exists())
        for name in ('ancestor', 'ours', 'theirs'):
            self.assertTrue((self.artifacts / (name + '.xlsm')).exists())


if __name__ == '__main__':
    unittest.main()




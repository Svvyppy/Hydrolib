import sys
from pathlib import Path
import tempfile
import unittest
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools/xlsm_merge'))
from file_refs import scan_file_refs, normalize
from validate import validate

class FileReferenceTests(unittest.TestCase):
    def book(self):
        w = openpyxl.Workbook()
        w.active.append(['Part Number', 'Library Path', 'Library Ref', 'Footprint Path 1', 'Footprint Ref 1'])
        return w

    def test_literals_and_formula_dependencies_resolve_without_cache(self):
        w = self.book()
        w.active.append(['A', '=CONCATENATE("Schematic\\",A2,".SchLib")', '=A2',
                         '="Footprints/"&E2&".PcbLib"', '=CONCAT("LQ","FP")'])
        self.assertEqual(scan_file_refs(w, {'Schematic/A.SchLib', 'Footprints/LQFP.PcbLib'}), [])

    def test_multiple_slots_and_optional_empty_formula(self):
        w = self.book()
        w.active.append(['A', 'Schematic/A.SchLib', 'A', '=CONCATENATE("Footprints/",E2,".PcbLib")', None])
        w.active.cell(1, 6, 'Footprint Path 2')
        w.active.cell(2, 6, 'Footprints/Missing.PcbLib')
        issues = scan_file_refs(w, {'Schematic/A.SchLib'})
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['cell'], 'F2')
        self.assertEqual(issues[0]['code'], 'missing_file')

    def test_case_path_escape_and_unsupported_formula(self):
        w = self.book()
        w.active.append(['A', 'schematic/a.schlib', 'A', '../outside.PcbLib', 'x'])
        w.active.append(['B', '=WEBSERVICE("https://example.com")', 'B'])
        issues = scan_file_refs(w, {'Schematic/A.SchLib'})
        self.assertEqual([i['code'] for i in issues], ['file_case', 'unchecked_path', 'invalid_path'])
        self.assertEqual(normalize('./Footprints\\x.PcbLib'), 'Footprints/x.PcbLib')
        for path in ('C:\\x.SchLib', '//server/x', '../../x'):
            with self.assertRaises(ValueError):
                normalize(path)

    def test_cycle_is_explicitly_unchecked(self):
        w = self.book()
        w.active.append(['A', '=C2', '=B2'])
        self.assertEqual(scan_file_refs(w, set())[0]['code'], 'unchecked_path')

    def test_removed_file_blocks_but_legacy_missing_warns(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            w = self.book()
            w.active.append(['A', 'Schematic/A.SchLib', 'A', 'Footprints/OLD.PcbLib', 'OLD'])
            path = d / 'w.xlsx'
            w.save(path)
            report = validate(path, path, d / 'validation.json', 'PR', set(), {'Schematic/A.SchLib'})
            self.assertEqual(report['errors'], 1)
            self.assertEqual(report['warnings'], 1)
            self.assertIn('Schematic/A.SchLib', next(i['message'] for i in report['issues'] if i['severity'] == 'error'))
            report = validate(path, path, d / 'validation.json', 'PR', {'Schematic/A.SchLib', 'Footprints/OLD.PcbLib'}, {'Schematic/A.SchLib'})
            self.assertEqual(report['errors'], 0)
            self.assertEqual(report['warnings'], 0)


    def test_schematic_formula_is_checked_even_when_library_ref_is_empty(self):
        w = self.book()
        w.active.append(['A', '=CONCATENATE("Schematic/",A2,".SchLib")'])
        self.assertEqual(scan_file_refs(w, set())[0]['code'], 'missing_file')

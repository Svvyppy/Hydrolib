from pathlib import Path
import sys
import json
import subprocess
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as E
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools/xlsm_merge'))
from merge import merge, Conflict, Q, Resolver


class MergeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.base = self.root / 'base.xlsx'
        w = openpyxl.Workbook()
        s = w.active
        s.title = 'Parts'
        s.append(['Part Number', 'Value', 'Name'])
        s.append(['A', 10, '=A2&B2'])
        w.save(self.base)
        self.ours = self.variant('ours', {})
        self.theirs = self.variant('theirs', {})
        self.output = self.root / 'output.xlsx'

    def variant(self, name, edits, delete=False, part=None):
        with ZipFile(self.base) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        r = E.fromstring(parts['xl/worksheets/sheet1.xml'])
        data = r.find(Q('sheetData'))
        if delete:
            data.remove(data[1])
        for coord, value in edits.items():
            rownum = ''.join(c for c in coord if c.isdigit())
            row = data.find(f"{Q('row')}[@r='{rownum}']")
            if row is None:
                row = E.SubElement(data, Q('row'), r=rownum)
            c = row.find(f"{Q('c')}[@r='{coord}']")
            if c is not None:
                row.remove(c)
            c = E.SubElement(row, Q('c'), r=coord)
            if isinstance(value, str) and value.startswith('='):
                E.SubElement(c, Q('f')).text = value[1:]
            elif isinstance(value, str):
                c.set('t', 'inlineStr')
                E.SubElement(E.SubElement(c, Q('is')), Q('t')).text = value
            else:
                E.SubElement(c, Q('v')).text = str(value)
        parts['xl/worksheets/sheet1.xml'] = E.tostring(r)
        if part:
            parts[part[0]] = part[1]
        p = self.root / (name + '.xlsx')
        with ZipFile(p, 'w', ZIP_DEFLATED) as z:
            for n, value in parts.items():
                z.writestr(n, value)
        return p

    def run_merge(self):
        merge(self.base, self.ours, self.theirs, self.output)
        return openpyxl.load_workbook(self.output)['Parts']

    def test_independent_cells(self):
        self.ours = self.variant('ours', {'B2': 20})
        self.theirs = self.variant('theirs', {'C2': '=A2&"x"'})
        s = self.run_merge()
        self.assertEqual(s['B2'].value, 20)
        self.assertEqual(s['C2'].value, '=A2&"x"')

    def test_concurrent_append_translates_formula(self):
        self.ours = self.variant('ours', {'A3': 'B', 'B3': 2, 'C3': '=A3&B3'})
        self.theirs = self.variant('theirs', {'A3': 'C', 'B3': 3, 'C3': '=A3&B3'})
        s = self.run_merge()
        self.assertEqual(s['A3'].value, 'B')
        self.assertEqual(s['A4'].value, 'C')
        self.assertEqual(s['C4'].value, '=A4&B4')

    def test_conflict_is_transactional(self):
        self.ours = self.variant('ours', {'B2': 20})
        self.theirs = self.variant('theirs', {'B2': 30})
        original = self.ours.read_bytes()
        with self.assertRaises(Conflict):
            merge(self.base, self.ours, self.theirs, self.ours)
        self.assertEqual(self.ours.read_bytes(), original)

    def test_delete_edit_conflicts(self):
        self.ours = self.variant('ours', {}, delete=True)
        self.theirs = self.variant('theirs', {'B2': 20})
        with self.assertRaises(Conflict):
            self.run_merge()

    def test_delete_unchanged(self):
        self.ours = self.variant('ours', {}, delete=True)
        self.assertIsNone(self.run_merge()['A2'].value)

    def test_duplicate_identity(self):
        self.theirs = self.variant('theirs', {'A3': 'A'})
        with self.assertRaises(Conflict):
            self.run_merge()

    def test_cross_row_reference_cannot_move(self):
        self.ours = self.variant('ours', {'A3': 'B'})
        self.theirs = self.variant('theirs', {'A3': 'C', 'C3': '=B2'})
        with self.assertRaises(Conflict):
            self.run_merge()

    def test_cross_row_reference_unchanged(self):
        self.base = self.variant('newbase', {'C2': '=B99'})
        self.ours = self.variant('ours', {})
        self.theirs = self.variant('theirs', {})
        self.assertEqual(self.run_merge()['C2'].value, '=B99')

    def discover(self):
        resolver = Resolver((self.base, self.ours, self.theirs))
        with self.assertRaises(Conflict):
            merge(self.base, self.ours, self.theirs, self.output, resolver)
        self.assertFalse(self.output.exists())
        return resolver.report('conflict')

    def resolved(self, document):
        resolver = Resolver((self.base, self.ours, self.theirs), document)
        merge(self.base, self.ours, self.theirs, self.output, resolver)
        return openpyxl.load_workbook(self.output)['Parts']

    def test_collect_and_resolve_all_with_independent_edits(self):
        self.ours = self.variant('ours', {'B2': 20, 'C2': '=A2&"ours"', 'A3': 'B', 'B3': 9})
        self.theirs = self.variant('theirs', {'B2': 30, 'C2': '=A2&"theirs"', 'A3': 'C', 'B3': 8})
        doc = self.discover()
        self.assertEqual(len(doc['conflicts']), 2)
        self.assertEqual(doc['conflicts'][0]['field'], 'Value')
        for c in doc['conflicts']:
            c['choice'] = 'ours' if c['column'] == 'B' else 'theirs'
        s = self.resolved(doc)
        self.assertEqual(s['B2'].value, 20)
        self.assertEqual(s['C2'].value, '=A2&"theirs"')
        self.assertEqual(s['A3'].value, 'B')
        self.assertEqual(s['A4'].value, 'C')
        self.assertEqual(s['B4'].value, 8)

    def test_delete_edit_selects_whole_record(self):
        self.ours = self.variant('ours', {}, delete=True)
        self.theirs = self.variant('theirs', {'B2': 30})
        doc = self.discover()
        self.assertEqual(doc['conflicts'][0]['kind'], 'record')
        self.assertEqual(doc['conflicts'][0]['ours'], {'missing': True})
        doc['conflicts'][0]['choice'] = 'theirs'
        self.assertEqual(self.resolved(doc)['B2'].value, 30)
        doc['conflicts'][0]['choice'] = 'ours'
        self.assertIsNone(self.resolved(doc)['A2'].value)

    def test_conflicting_new_identity_chooses_whole_record(self):
        self.ours = self.variant('ours', {'A3': 'NEW', 'B3': 10})
        self.theirs = self.variant('theirs', {'A3': 'NEW', 'C3': 'from theirs'})
        doc = self.discover()
        doc['conflicts'][0]['choice'] = 'theirs'
        s = self.resolved(doc)
        self.assertIsNone(s['B3'].value)
        self.assertEqual(s['C3'].value, 'from theirs')

    def test_incomplete_choices_leave_output_untouched(self):
        self.ours = self.variant('ours', {'B2': 20, 'C2': 'ours'})
        self.theirs = self.variant('theirs', {'B2': 30, 'C2': 'theirs'})
        doc = self.discover()
        doc['conflicts'][0]['choice'] = 'ours'
        self.output.write_bytes(b'original output')
        with self.assertRaises(Conflict):
            self.resolved(doc)
        self.assertEqual(self.output.read_bytes(), b'original output')

    def test_stale_choices_rejected(self):
        self.ours = self.variant('ours', {'B2': 20})
        self.theirs = self.variant('theirs', {'B2': 30})
        doc = self.discover()
        doc['conflicts'][0]['choice'] = 'ours'
        self.theirs = self.variant('theirs', {'B2': 40})
        with self.assertRaisesRegex(Conflict, 'stale'):
            self.resolved(doc)

    def test_invalid_and_unknown_choices_rejected(self):
        self.ours = self.variant('ours', {'B2': 20})
        self.theirs = self.variant('theirs', {'B2': 30})
        doc = self.discover()
        doc['conflicts'][0]['choice'] = 'base'
        with self.assertRaisesRegex(Conflict, 'choice must'):
            self.resolved(doc)
        doc['conflicts'][0]['choice'] = 'ours'
        doc['conflicts'][0]['id'] = 'unknown'
        with self.assertRaisesRegex(Conflict, 'unknown'):
            self.resolved(doc)

    def test_cli_report_can_be_reused_as_resolutions(self):
        self.ours = self.variant('ours', {'B2': 20})
        self.theirs = self.variant('theirs', {'B2': 30})
        script = Path(__file__).resolve().parents[1] / 'tools/xlsm_merge/merge.py'
        report = self.root / 'report.json'
        command = [sys.executable, str(script), str(self.base), str(self.ours),
                   str(self.theirs), str(self.output), '--report', str(report)]
        self.assertEqual(subprocess.run(command, capture_output=True).returncode, 1)
        doc = json.loads(report.read_text())
        doc['conflicts'][0]['choice'] = 'theirs'
        decisions = self.root / 'resolutions.json'
        decisions.write_text(json.dumps(doc))
        self.assertEqual(subprocess.run(command + ['--resolutions', str(decisions)], capture_output=True).returncode, 0)
        self.assertEqual(openpyxl.load_workbook(self.output)['Parts']['B2'].value, 30)
        self.assertEqual(json.loads(report.read_text())['status'], 'merged')

    def test_default_side_preserves_both_independent_changes(self):
        self.ours = self.variant('ours', {'B2': 20, 'A3': 'B', 'B3': 9})
        self.theirs = self.variant('theirs', {'B2': 30, 'C2': 'independent'})
        for side, expected in [('ours', 20), ('theirs', 30)]:
            resolver = Resolver((self.base, self.ours, self.theirs), conflict_side=side)
            merge(self.base, self.ours, self.theirs, self.output, resolver)
            s = openpyxl.load_workbook(self.output)['Parts']
            self.assertEqual(s['B2'].value, expected)
            self.assertEqual(s['C2'].value, 'independent')
            self.assertEqual(s['A3'].value, 'B')
            self.assertEqual(resolver.report('merged')['conflicts'][0]['choice'], side)

    def test_explicit_choice_overrides_default_and_null_stays_unresolved(self):
        self.ours = self.variant('ours', {'B2': 20})
        self.theirs = self.variant('theirs', {'B2': 30})
        doc = self.discover()
        doc['conflicts'][0]['choice'] = 'theirs'
        resolver = Resolver((self.base, self.ours, self.theirs), doc, conflict_side='ours')
        merge(self.base, self.ours, self.theirs, self.output, resolver)
        self.assertEqual(openpyxl.load_workbook(self.output)['Parts']['B2'].value, 30)
        doc['conflicts'][0]['choice'] = None
        resolver = Resolver((self.base, self.ours, self.theirs), doc, conflict_side='ours')
        with self.assertRaises(Conflict):
            merge(self.base, self.ours, self.theirs, self.output, resolver)

    def test_default_side_cannot_override_schema_guard(self):
        self.theirs = self.variant('theirs', {'B1': 'New schema'})
        resolver = Resolver((self.base, self.ours, self.theirs), conflict_side='ours')
        with self.assertRaisesRegex(Conflict, 'schema'):
            merge(self.base, self.ours, self.theirs, self.output, resolver)
        self.assertFalse(self.output.exists())

    def test_real_library_parallel_additions(self):
        self.base = Path(__file__).resolve().parents[1] / 'Database/Hydrolib_bd.xlsm'
        n = openpyxl.load_workbook(self.base)['Chips'].max_row + 1
        self.ours = self.variant('ours', {f'A{n}': 'TEST-OURS', f'C{n}': f'=A{n}'})
        self.theirs = self.variant('theirs', {f'A{n}': 'TEST-THEIRS', f'C{n}': f'=A{n}'})
        merge(self.base, self.ours, self.theirs, self.output)
        s = openpyxl.load_workbook(self.output)['Chips']
        self.assertEqual(s[f'A{n}'].value, 'TEST-OURS')
        self.assertEqual(s[f'A{n+1}'].value, 'TEST-THEIRS')
        self.assertEqual(s[f'C{n+1}'].value, f'=A{n+1}')
        self.assertEqual(next(iter(s.tables.values())).ref, f'A1:M{n+1}')

    def test_real_library_roundtrip(self):
        actual = Path(__file__).resolve().parents[1] / 'Database/Hydrolib_bd.xlsm'
        merge(actual, actual, actual, self.output)
        a = openpyxl.load_workbook(actual)
        b = openpyxl.load_workbook(self.output)
        for sheet in a:
            for row in sheet:
                for cell in row:
                    self.assertEqual(cell.value, b[sheet.title][cell.coordinate].value)
        with ZipFile(actual) as x, ZipFile(self.output) as y:
            for n in x.namelist():
                if n not in {'xl/workbook.xml', 'xl/calcChain.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml'} and not n.startswith(('xl/worksheets/sheet', 'xl/tables/')):
                    self.assertEqual(x.read(n), y.read(n), n)


if __name__ == '__main__':
    unittest.main()

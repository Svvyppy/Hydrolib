import json
from pathlib import Path
import sys
import tempfile
import unittest
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools/xlsm_merge'))
from validate import validate, require_valid
from summary import render

class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.target = self.write('target', [['A', 10], [None, 'legacy']])

    def write(self, name, rows):
        path = self.root / (name + '.xlsx')
        w = openpyxl.Workbook()
        w.active.title = 'Parts'
        w.active.append(['Part Number', 'Value'])
        for row in rows:
            w.active.append(row)
        w.save(path)
        w.close()
        return path

    def run_check(self, path):
        return validate(path, self.target, self.root / 'validation.json', 'PR')

    def test_legacy_blank_warns_new_blank_blocks(self):
        self.assertEqual(self.run_check(self.target)['errors'], 0)
        candidate = self.write('head', [['A', 10], [None, 'legacy'], [' ', 'new']])
        result = self.run_check(candidate)
        self.assertEqual(result['errors'], 1)
        self.assertEqual(result['warnings'], 1)
        with self.assertRaisesRegex(ValueError, 'Проверка компонентов'):
            require_valid(candidate, self.target, self.root / 'validation.json', 'PR')

    def test_duplicate_formula_and_excel_error(self):
        candidate = self.write('head', [['A', 10], ['A', '#REF!'], ['=A2', 10]])
        result = self.run_check(candidate)
        self.assertEqual(result['errors'], 3)
        self.assertEqual({i['code'] for i in result['issues']}, {'duplicate_part', 'formula_part', 'excel_error'})

    def test_whitespace_warning_and_independent_sheets(self):
        candidate = self.write('head', [['A', 10], [' B ', 10]])
        w = openpyxl.load_workbook(candidate)
        s = w.create_sheet('Other')
        s.append(['Part Number', 'Value'])
        s.append(['A', 20])
        w.save(candidate)
        w.close()
        result = self.run_check(candidate)
        self.assertEqual(result['errors'], 0)
        self.assertEqual(result['warnings'], 1)

    def test_unchanged_excel_error_warns_and_formula_is_not_evaluated(self):
        target = self.write('old', [['A', '#REF!'], ['B', '=1/0']])
        result = validate(target, target, self.root / 'validation.json', 'result')
        self.assertEqual(result['errors'], 0)
        self.assertEqual(result['warnings'], 1)

    def test_report_escapes_values_and_explains_failure(self):
        self.run_check(self.write('head', [['<script>', '#REF!']]))
        (self.root / 'failure.txt').write_text('Проверка компонентов')
        result = render(self.root, 'failure', 'a/b', '1')
        self.assertIn('Ошибок: 1', result)
        self.assertIn('&lt;script&gt;', result)
        self.assertIn('B2', result)
        self.assertIn('validation.json', result)
        self.assertNotIn('Если причина — структура', result)

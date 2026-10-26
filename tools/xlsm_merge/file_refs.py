"""Resolve a small, data-only subset of Excel path formulas against Git files."""
import re
from pathlib import PurePosixPath
from openpyxl.formula.tokenizer import Tokenizer


class UnsupportedFormula(ValueError):
    pass


def value(sheet, coordinate, stack=()):
    coordinate = coordinate.replace('$', '').upper()
    if coordinate in stack or len(stack) >= 32:
        raise UnsupportedFormula('циклическая или слишком длинная цепочка ссылок')
    cell = sheet[coordinate]
    if cell.data_type == 'e':
        raise UnsupportedFormula('ошибка Excel в зависимости пути')
    if cell.data_type != 'f':
        return '' if cell.value is None else str(cell.value)
    tokens = [t for t in Tokenizer(cell.value).items if t.type != 'WHITE-SPACE']
    pos = 0
    def atom():
        nonlocal pos
        if pos >= len(tokens):
            raise UnsupportedFormula('неполная формула')
        token = tokens[pos]
        pos += 1
        if token.type == 'OPERAND':
            if token.subtype == 'TEXT':
                return token.value[1:-1].replace('""', '"')
            if token.subtype == 'NUMBER':
                return token.value
            if token.subtype == 'RANGE' and re.fullmatch(r'\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}', token.value):
                return value(sheet, token.value, (*stack, coordinate))
        if token.type == 'FUNC' and token.subtype == 'OPEN' and token.value.upper() in ('CONCATENATE(', 'CONCAT('):
            args = []
            while True:
                args.append(expr())
                if pos >= len(tokens):
                    break
                t = tokens[pos]
                pos += 1
                if t.type == 'FUNC' and t.subtype == 'CLOSE':
                    return ''.join(args)
                if t.type != 'SEP' or t.subtype != 'ARG':
                    break
        raise UnsupportedFormula('поддерживаются текст, ссылки на ячейки, CONCATENATE/CONCAT и &')
    def expr():
        nonlocal pos
        result = atom()
        while pos < len(tokens) and tokens[pos].type == 'OPERATOR-INFIX' and tokens[pos].value == '&':
            pos += 1
            result += atom()
        return result
    result = expr()
    if pos != len(tokens):
        raise UnsupportedFormula('неподдерживаемое выражение пути')
    return result


def normalize(path):
    path = path.replace('\\', '/')
    if path.startswith('/') or re.match(r'^[A-Za-z]:', path):
        raise ValueError('абсолютный путь; укажите путь от корня репозитория')
    pieces = []
    for part in path.split('/'):
        if part in ('', '.'):
            continue
        if part == '..':
            if not pieces:
                raise ValueError('путь выходит за корень репозитория')
            pieces.pop()
        else:
            pieces.append(part)
    return str(PurePosixPath(*pieces))


def scan_file_refs(book, files):
    files = set(files)
    folded = {}
    for path in sorted(files):
        folded.setdefault(path.casefold(), []).append(path)
    issues = []
    for sheet in book:
        headers = {str(c.value): c.column for c in sheet[1] if c.value is not None}
        for heading, col in headers.items():
            if heading != 'Library Path' and not re.fullmatch(r'Footprint Path \d+', heading):
                continue
            ref_heading = heading.replace('Path', 'Ref')
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, col)
                if cell.value is None:
                    continue
                def add(code, message):
                    issues.append({'sheet': sheet.title, 'row': row,
                        'part': str(sheet.cell(row, 1).value or ''), 'cell': cell.coordinate,
                        'field': heading, 'code': code, 'message': message})
                try:
                    # Optional footprint formulas are often pre-filled for empty Ref slots.
                    if heading.startswith('Footprint Path ') and cell.data_type == 'f' and ref_heading in headers:
                        if not value(sheet, sheet.cell(row, headers[ref_heading]).coordinate):
                            continue
                    path = value(sheet, cell.coordinate)
                    if not path:
                        continue
                    path = normalize(path)
                    if path in files:
                        continue
                    matches = folded.get(path.casefold(), [])
                    if len(matches) == 1:
                        add('file_case', f'Регистр пути отличается: {path}; файл в Git: {matches[0]}.')
                    else:
                        add('missing_file', f'Файл не найден в Git: {path}. Добавьте файл или исправьте путь.')
                except UnsupportedFormula as error:
                    add('unchecked_path', f'Путь не проверен: {error}.')
                except ValueError as error:
                    add('invalid_path', f'Некорректный путь: {error}.')
    return issues

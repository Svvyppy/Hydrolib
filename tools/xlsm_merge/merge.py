"""Conservative, transactional three-way merge of Hydrolib OOXML workbooks.

Read formulas with openpyxl; write original ZIP parts with lxml, never save via
openpyxl (that would discard XML maps). Unsupported changes fail closed.
"""
import argparse
import hashlib
from datetime import date, datetime, time
from copy import deepcopy
from io import BytesIO
import json
import os
import tempfile
from pathlib import Path
import re
import sys
import zipfile

from lxml import etree as ET
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
Q = lambda tag: '{' + NS + '}' + tag
MISSING = object()


class Conflict(ValueError):
    pass


def choose(base, ours, theirs, location):
    if ours == theirs or theirs == base:
        return ours
    if ours == base:
        return theirs
    raise Conflict(f'{location}: conflicting changes; base={readable(base)!r}; ours={readable(ours)!r}; theirs={readable(theirs)!r}')



def readable(value):
    """JSON representation without unstable object reprs or executable payloads."""
    if value is MISSING:
        return {'missing': True}
    if isinstance(value, (datetime, date, time)):
        return {'type': type(value).__name__, 'value': value.isoformat()}
    if isinstance(value, bytes):
        return {'sha256': hashlib.sha256(value).hexdigest()}
    if isinstance(value, dict):
        return {str(k): readable(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [readable(v) for v in value]
    return value


class Resolver:
    """Collect all data conflicts and apply choices bound to exact input bytes."""
    def __init__(self, paths, document=None, conflict_side='manual'):
        if conflict_side not in ('manual', 'ours', 'theirs'):
            raise Conflict('Conflict side must be manual, ours or theirs')
        self.conflict_side = None if conflict_side == 'manual' else conflict_side
        self.inputs = {name: hashlib.sha256(Path(path).read_bytes()).hexdigest()
                       for name, path in zip(('base', 'ours', 'theirs'), paths)}
        self.choices = {}
        self.conflicts = []
        self.seen = set()
        if document is not None:
            if document.get('version') != 1 or document.get('inputs') != self.inputs:
                raise Conflict('Resolution file is stale: regenerate it for these three workbook versions')
            records = document.get('conflicts')
            if not isinstance(records, list):
                raise Conflict('Resolution file must contain a conflicts array')
            for record in records:
                identity, choice = record.get('id'), record.get('choice')
                if not isinstance(identity, str) or identity in self.choices:
                    raise Conflict('Invalid or duplicate conflict id')
                if choice not in (None, 'ours', 'theirs'):
                    raise Conflict(f'{identity}: choice must be ours, theirs or null')
                self.choices[identity] = choice

    def choose(self, base, ours, theirs, context, display=None):
        try:
            return choose(base, ours, theirs, str(context))
        except Conflict:
            pass
        identity = hashlib.sha256(json.dumps(context, ensure_ascii=False,
                                             sort_keys=True).encode()).hexdigest()[:24]
        self.seen.add(identity)
        choice = self.choices.get(identity, self.conflict_side)
        values = display if display is not None else (base, ours, theirs)
        self.conflicts.append({'id': identity, **context,
                               **dict(zip(('base', 'ours', 'theirs'), map(readable, values))),
                               'choice': choice})
        # Temporary fallback is only for discovery. No workbook is written until
        # every conflict has an explicit choice and all safety checks pass.
        return theirs if choice == 'theirs' else ours

    def finish(self):
        if set(self.choices) - self.seen:
            raise Conflict('Resolution file contains unknown conflict ids')
        unresolved = sum(c['choice'] is None for c in self.conflicts)
        if unresolved:
            raise Conflict(f'{unresolved} unresolved conflict(s); set each choice to ours or theirs')

    def report(self, status, message=None):
        result = {'version': 1, 'status': status, 'inputs': self.inputs,
                  'conflicts': self.conflicts, 'conflict_side': self.conflict_side or 'manual'}
        if message:
            result['message'] = message
        return result


def formula(value, origin, target):
    if origin == target:
        return value
    row = int(re.search(r'\d+$', origin)[0])
    for token in Tokenizer(value).items:
        if token.type != 'OPERAND' or token.subtype != 'RANGE':
            continue
        ref = token.value
        if ('[#This Row]' in ref or '[@' in ref) and '!' not in ref:
            continue
        # Only references to this row are safe to relocate. External references,
        # ranges of other rows, named ranges, and absolute rows fail closed.
        if not re.fullmatch(r'\$?[A-Za-z]+\d+(?::\$?[A-Za-z]+\d+)?', ref):
            raise Conflict(f'{origin}: unsupported reference {ref!r}')
        if any(int(n) != row for n in re.findall(r'\d+', ref)):
            raise Conflict(f'{origin}: reference to another row: {ref}')
    return Translator(value, origin=origin).translate_formula(target)


def style_signature(element):
    """Namespace-prefix-independent XML key; normalize equivalent booleans."""
    def value(text):
        return {'true': '1', 'false': '0'}.get(text, text)
    return (element.tag, tuple(sorted((k, value(v)) for k, v in element.attrib.items())),
            (element.text or '').strip(), tuple(style_signature(c) for c in element))


class StylePool:
    """Reconcile cell-XF indices while keeping shared style dependencies fixed."""
    def __init__(self, books):
        roots = [ET.fromstring(b.parts['xl/styles.xml']) for b in books]
        def dependencies(root):
            return tuple(style_signature(c) for c in root if c.tag != Q('cellXfs'))
        if len({dependencies(r) for r in roots}) != 1:
            raise Conflict('Style dependencies changed (fonts/fills/borders/number formats); manual review required')
        root = deepcopy(roots[1])
        target = root.find(Q('cellXfs'))
        if target is None or any(c.tag != Q('xf') for c in target):
            raise Conflict('Unsupported cell style table')
        # Keep all ours indices unchanged for existing worksheet metadata.
        indices = {style_signature(c): i for i, c in enumerate(target)}
        maps = [None, {i: i for i in range(len(target))}, None]
        for side in (0, 2):
            mapping = {}
            for i, cell_style in enumerate(roots[side].find(Q('cellXfs'))):
                if cell_style.tag != Q('xf'):
                    raise Conflict('Unsupported cell style record')
                key = style_signature(cell_style)
                if key not in indices:
                    indices[key] = len(target)
                    target.append(deepcopy(cell_style))
                mapping[i] = indices[key]
            maps[side] = mapping
        # Duplicate ours entries must compare equal as well as display equally.
        maps[1] = {i: indices[style_signature(c)] for i, c in enumerate(roots[1].find(Q('cellXfs')))}
        for book, mapping in zip(books, maps):
            book.style_map = mapping
        target.set('count', str(len(target)))
        self.data = (books[1].parts['xl/styles.xml']
                     if all(b.parts['xl/styles.xml'] == books[1].parts['xl/styles.xml'] for b in books)
                     else xml(root))


class Book:
    def __init__(self, path):
        self.raw = Path(path).read_bytes()
        with zipfile.ZipFile(BytesIO(self.raw)) as z:
            self.parts = {n: z.read(n) for n in z.namelist()}
        self.wb = openpyxl.load_workbook(BytesIO(self.raw), keep_vba=False, data_only=False)
        root = ET.fromstring(self.parts['xl/workbook.xml'])
        rels = ET.fromstring(self.parts['xl/_rels/workbook.xml.rels'])
        targets = {r.get('Id'): r.get('Target') for r in rels}
        self.sheets = {}
        for s in root.find(Q('sheets')):
            target = targets[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
            self.sheets[s.get('name')] = target.lstrip('/') if target.startswith('/') else 'xl/' + target
        self.cell_xml = {}
        for sheet, path in self.sheets.items():
            root = ET.fromstring(self.parts[path])
            self.cell_xml[sheet] = {c.get('r'): c for row in root.find(Q('sheetData')) for c in row.findall(Q('c'))}
        self.strings = []
        if 'xl/sharedStrings.xml' in self.parts:
            for si in ET.fromstring(self.parts['xl/sharedStrings.xml']):
                self.strings.append(si)

    def rows(self, name):
        ws = self.wb[name]
        root = ET.fromstring(self.parts[self.sheets[name]])
        rows = {}
        order = []
        for row in root.find(Q('sheetData')):
            number = int(row.get('r'))
            cells = {c.get('r'): c for c in row.findall(Q('c'))}
            meaningful = [c for c in ws[number] if c.value is not None]
            if number == 1 or not meaningful:
                key = ('fixed', number)
            else:
                identity = ws.cell(number, 1)
                if identity.data_type == 'f':
                    raise Conflict(f'{name}!A{number}: formula cannot be a component identity')
                key = ('part', str(identity.value)) if identity.value is not None else ('unkeyed', number)
            if key in rows:
                raise Conflict(f'{name}: duplicate Part Number {key[1]!r}')
            values = {}
            for coord, c in cells.items():
                cell = ws[coord]
                value = cell.value
                if cell.data_type == 'f':
                    if not isinstance(value, str):
                        raise Conflict(f'{name}!{coord}: array/data-table formula is unsupported')
                    try:
                        value = formula(value, coord, f'{cell.column_letter}1')
                    except Conflict:
                        value = ('anchored', coord, value)
                # Compare content and cell style; cached formula values are ignored.
                values[cell.column] = (cell.data_type, value, self.style_map[int(c.get('s', '0'))])
            rows[key] = (number, row, values)
            order.append(key)
        return root, rows, order

    def copy_cell(self, name, col, row, dest):
        coord = f'{get_column_letter(col)}{row}'
        src = self.cell_xml[name].get(coord)
        if src is None:
            return None
        c = deepcopy(src)
        target = f'{get_column_letter(col)}{dest}'
        c.set('r', target)
        mapped_style = self.style_map[int(c.get('s', '0'))]
        if c.get('s') is not None or mapped_style:
            c.set('s', str(mapped_style))
        if c.get('t') == 's':
            si = self.strings[int(c.find(Q('v')).text)]
            c.remove(c.find(Q('v')))
            c.set('t', 'inlineStr')
            inline = ET.SubElement(c, Q('is'))
            for child in si:
                inline.append(deepcopy(child))
        f = c.find(Q('f'))
        if f is not None:
            # Expand shared formulas before moving them. Drop stale cached result.
            value = self.wb[name][coord].value
            f.attrib.clear()
            f.text = formula(value, coord, target)[1:]
            v = c.find(Q('v'))
            if v is not None:
                c.remove(v)
        return c


def xml(data):
    return ET.tostring(data, xml_declaration=True, encoding='UTF-8', standalone=True)


def merge(base_path, ours_path, theirs_path, output, resolver=None):
    resolver = resolver or Resolver((base_path, ours_path, theirs_path))
    books = [Book(p) for p in (base_path, ours_path, theirs_path)]
    base, ours, theirs = books
    style_pool = StylePool(books)
    if not (base.sheets == ours.sheets == theirs.sheets):
        raise Conflict('Sheet names/order/relationships changed; manual review required')
    result = dict(ours.parts)
    result['xl/styles.xml'] = style_pool.data
    handled = set(ours.sheets.values()) | {'xl/styles.xml'}
    # Shared string indices are local to each ZIP; all emitted cells use inline strings.
    handled.update({'xl/sharedStrings.xml', 'xl/calcChain.xml', 'docProps/core.xml', 'docProps/app.xml'})
    tables = {n for b in books for n in b.parts if n.startswith('xl/tables/') and n.endswith('.xml')}
    handled |= tables
    # Preserve all package features. Never silently discard a macro or XML map.
    for name in sorted(set().union(*(b.parts for b in books)) - handled):
        values = [b.parts.get(name, MISSING) for b in books]
        if name == 'xl/workbook.xml':
            normalized = []
            for v in values:
                r = ET.fromstring(v)
                for child in list(r):
                    if child.tag in {Q('calcPr'), Q('bookViews')}:
                        r.remove(child)
                normalized.append(ET.tostring(r, method='c14n'))
            if len(set(normalized)) != 1:
                raise Conflict('Workbook structure changed')
            continue
        if name == 'xl/vbaProject.bin':
            value = resolver.choose(*values, {'kind': 'binary_part', 'part': name})
        else:
            value = choose(*values, name)
        if value is MISSING:
            result.pop(name, None)
        else:
            result[name] = value
    lengths = {}
    for name, path in ours.sheets.items():
        parsed = [b.rows(name) for b in books]
        # Ignore viewport and used-range changes; other sheet metadata (including
        # hyperlink anchors/relationships) must be identical to avoid broken links.
        def metadata(root, book):
            r = deepcopy(root)
            for child in list(r):
                if child.tag in {Q('sheetData'), Q('dimension'), Q('sheetViews'), Q('cols')}:
                    r.remove(child)
            for col in r.iter(Q('col')):
                if col.get('style') is not None:
                    col.set('style', str(book.style_map[int(col.get('style'))]))
            return style_signature(r)
        if len({metadata(p[0], book) for p, book in zip(parsed, books)}) != 1:
            raise Conflict(f'{name}: worksheet metadata changed; manual review required')
        root = deepcopy(parsed[1][0])
        column_nodes = []
        for parsed_sheet, book in zip(parsed, books):
            cols = deepcopy(parsed_sheet[0].find(Q('cols')))
            if cols is not None:
                for col in cols:
                    if col.get('style') is not None:
                        col.set('style', str(book.style_map[int(col.get('style'))]))
            column_nodes.append(cols)
        signatures = [style_signature(c) if c is not None else MISSING for c in column_nodes]
        selected = resolver.choose(*signatures, {'kind': 'column_layout', 'sheet': name})
        existing = root.find(Q('cols'))
        position = list(root).index(existing) if existing is not None else list(root).index(root.find(Q('sheetData')))
        if existing is not None:
            root.remove(existing)
        if selected is not MISSING:
            source = next(i for i in (1, 2, 0) if signatures[i] == selected)
            root.insert(position, column_nodes[source])
        data = root.find(Q('sheetData'))
        data.clear()
        maps = [p[1] for p in parsed]
        headers = [m.get(('fixed', 1)) for m in maps]
        schemas = [{col: value[:2] for col, value in h[2].items() if value[1] is not None}
                   if h else None for h in headers]
        if any(h is None for h in headers) or not (schemas[0] == schemas[1] == schemas[2]):
            raise Conflict(f'{name}: column schema changed')
        for key, record in maps[0].items():
            for mapping in maps[1:]:
                if key in mapping and mapping[key][0] != record[0]:
                    raise Conflict(f'{name}/{key[1]}: existing rows moved; review references manually')
        # Preserve our physical order; append new records from the other branch.
        order = parsed[1][2] + [k for k in parsed[2][2] if k not in maps[1]]
        for key in maps[0]:
            if key not in order:
                order.append(key)
        last = 0
        for key in order:
            records = [m.get(key) for m in maps]
            signatures = [r[2] if r else MISSING for r in records]
            context = {'sheet': name, 'record_type': key[0], 'part': key[1]}
            # A delete/edit conflict selects the whole record. After selection,
            # exclude other versions so cell merging cannot resurrect old fields.
            if any(r is None for r in records) or key[0] in {'fixed', 'unkeyed'}:
                merged = resolver.choose(*signatures, {**context, 'kind': 'record'})
                if merged is MISSING:
                    continue
                selected = next(i for i in (1, 2, 0) if signatures[i] == merged)
                records = [r if i == selected else None for i, r in enumerate(records)]
            # Keep existing gaps to preserve static anchors; imported rows append.
            dest = max(last + 1, records[1][0] if records[1] else last + 1)
            if key[0] != 'part' and any(r and r[0] != dest for r in records):
                raise Conflict(f'{name}/{key}: cannot move an unkeyed or fixed row')
            source = next(i for i in (1, 2, 0) if records[i])
            attributes = [dict(r[1].attrib) if r else MISSING for r in records]
            for side, attr in enumerate(attributes):
                if attr is not MISSING:
                    attr.pop('r', None)
                    attr.pop('spans', None)
                    if 's' in attr:
                        attr['s'] = str(books[side].style_map[int(attr['s'])])
            selected_attrs = resolver.choose(*attributes, {**context, 'kind': 'row_format'})
            row = deepcopy(records[source][1])
            row.attrib.clear()
            row.attrib.update(selected_attrs)
            row.set('r', str(dest))
            for child in list(row):
                if child.tag == Q('c'):
                    row.remove(child)
            columns = sorted(set().union(*(r[2] for r in records if r)))
            for col in columns:
                vals = [r[2].get(col, MISSING) if r else MISSING for r in records]
                display = [books[i].wb[name].cell(r[0], col).value if r and col in r[2] else MISSING
                           for i, r in enumerate(records)]
                value = resolver.choose(*vals, {**context, 'kind': 'cell',
                    'column': get_column_letter(col), 'field': ours.wb[name].cell(1, col).value}, display)
                if value is MISSING:
                    continue
                i = next(i for i in (1, 2, 0) if vals[i] == value)
                cell = books[i].copy_cell(name, col, records[i][0], dest)
                row.append(cell)
            data.append(row)
            last = dest
        dim = root.find(Q('dimension'))
        if dim is not None:
            dim.set('ref', f'A1:{get_column_letter(ours.wb[name].max_column)}{last}')
        lengths[name] = last
        result[path] = xml(root)
    # Table IDs, columns, XML-map bindings must stay unchanged. Only resize ranges.
    for path in sorted(tables):
        roots = [ET.fromstring(b.parts[path]) for b in books]
        def table_signature(r):
            r = deepcopy(r)
            r.attrib.pop('ref', None)
            af = r.find(Q('autoFilter'))
            if af is not None:
                af.attrib.pop('ref', None)
            return ET.tostring(r, method='c14n')
        if len({table_signature(r) for r in roots}) != 1:
            raise Conflict(f'{path}: table schema changed')
        root = roots[1]
        table_name = root.get('name')
        sheet = next(s.title for s in ours.wb if table_name in s.tables)
        start, end = root.get('ref').split(':')
        end = re.sub(r'\d+$', str(lengths[sheet]), end)
        root.set('ref', start + ':' + end)
        af = root.find(Q('autoFilter'))
        if af is not None:
            af.set('ref', root.get('ref'))
        result[path] = xml(root)
    # Rebuild calculation dependencies in Excel, retaining package relationships.
    if 'xl/calcChain.xml' in result:
        result.pop('xl/calcChain.xml')
        for path in ('xl/_rels/workbook.xml.rels', '[Content_Types].xml'):
            root = ET.fromstring(result[path])
            for child in list(root):
                if child.get('Type', '').endswith('/calcChain') or child.get('PartName') == '/xl/calcChain.xml':
                    root.remove(child)
            result[path] = xml(root)
    workbook = ET.fromstring(result['xl/workbook.xml'])
    calc = workbook.find(Q('calcPr'))
    if calc is None:
        calc = ET.SubElement(workbook, Q('calcPr'))
    calc.set('fullCalcOnLoad', '1')
    calc.set('forceFullCalc', '1')
    result['xl/workbook.xml'] = xml(workbook)
    resolver.finish()
    buf = BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, value in result.items():
            z.writestr(name, value)
    # Validate the whole result before replacing output (which may be ours).
    check = openpyxl.load_workbook(BytesIO(buf.getvalue()), keep_vba=False)
    check.close()
    destination = Path(output)
    with tempfile.NamedTemporaryFile(dir=destination.parent, delete=False) as temp:
        temporary = Path(temp.name)
        try:
            temp.write(buf.getvalue())
            temp.flush()
            os.fsync(temp.fileno())
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise
    try:
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for arg in ('base', 'ours', 'theirs', 'output'):
        parser.add_argument(arg)
    parser.add_argument('--report', default='xlsm-merge-report.json')
    parser.add_argument('--conflict-side', choices=('manual', 'ours', 'theirs'), default='manual',
                        help='Side for data conflicts only; explicit JSON choices take precedence')
    parser.add_argument('--resolutions', help='JSON report with choice=ours/theirs for each conflict')
    args = parser.parse_args()
    resolver = None
    try:
        document = json.loads(Path(args.resolutions).read_text(encoding='utf-8')) if args.resolutions else None
        resolver = Resolver((args.base, args.ours, args.theirs), document, args.conflict_side)
        merge(args.base, args.ours, args.theirs, args.output, resolver=resolver)
    except Exception as error:
        report = resolver.report('conflict', str(error)) if resolver else {'status': 'error', 'message': str(error)}
        Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(report['message'], file=sys.stderr)
        return 1
    Path(args.report).write_text(json.dumps(resolver.report('merged'), ensure_ascii=False, indent=2), encoding='utf-8')
    return 0


if __name__ == '__main__':
    sys.exit(main())

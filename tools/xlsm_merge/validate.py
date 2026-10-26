"""Read-only component checks. Never evaluates formulas or saves workbooks."""
from collections import Counter
import json
from pathlib import Path
import openpyxl


def scan(path):
    book = openpyxl.load_workbook(path, data_only=False)
    issues = []
    try:
        for sheet in book:
            seen = {}
            for cells in sheet.iter_rows(min_row=2):
                if not any(c.value is not None for c in cells):
                    continue
                identity = cells[0]
                part = identity.value
                context = {'sheet': sheet.title, 'row': identity.row,
                           'part': str(part) if part is not None else ''}
                def add(code, cell, message):
                    issues.append({**context, 'code': code, 'cell': cell.coordinate,
                                   'field': str(sheet.cell(1, cell.column).value or cell.column_letter),
                                   'message': message})
                if part is None or (isinstance(part, str) and not part.strip()):
                    add('missing_part', identity, 'Пустой Part Number. Укажите постоянный идентификатор компонента.')
                elif identity.data_type == 'f':
                    add('formula_part', identity, 'Part Number не должен вычисляться формулой.')
                else:
                    key = str(part)
                    if key in seen:
                        add('duplicate_part', identity, f'Повтор Part Number; первая запись в строке {seen[key]}.')
                    seen[key] = seen.get(key, identity.row)
                    if isinstance(part, str) and part != part.strip():
                        add('part_spaces', identity, 'Пробелы в начале или конце Part Number; проверьте идентификатор.')
                for cell in cells:
                    if cell.data_type == 'e':
                        add('excel_error', cell, f'Ошибка Excel {cell.value}; исправьте значение или формулу.')
    finally:
        book.close()
    return issues


def fingerprint(issue):
    # Keyed components retain identity when their physical position changes.
    return (issue['code'], issue['sheet'], issue['part'] or issue['row'],
            issue['field'], issue['message'])


def validate(path, target, output, phase):
    baseline = Counter(fingerprint(i) for i in scan(target))
    issues = scan(path)
    for issue in issues:
        key = fingerprint(issue)
        existing = baseline[key] > 0
        if existing:
            baseline[key] -= 1
        issue['existing'] = existing
        issue['severity'] = ('error' if issue['code'] in {'duplicate_part', 'formula_part'}
                             else 'warning' if existing or issue['code'] == 'part_spaces' else 'error')
    result = {'phase': phase, 'issues': issues,
              'errors': sum(i['severity'] == 'error' for i in issues),
              'warnings': sum(i['severity'] == 'warning' for i in issues)}
    Path(output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return result


def require_valid(path, target, output, phase):
    result = validate(path, target, output, phase)
    if result['errors']:
        raise ValueError(f"Проверка компонентов: ошибок {result['errors']}, предупреждений {result['warnings']}. Исправьте ошибки в книге; подробности в отчёте PR и validation.json.")
    return result
